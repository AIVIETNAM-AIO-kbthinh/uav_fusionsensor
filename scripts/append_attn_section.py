"""Chèn phần báo cáo về attention (F2b, F2c) vào data_report.docx.

    python scripts/append_attn_section.py

Idempotent: lần chạy đầu sao lưu bản gốc ra `data_report.base.docx`, mọi lần chạy
sau đều dựng lại từ bản sao lưu đó nên chạy bao nhiêu lần cũng không nhân đôi nội
dung. Muốn sửa câu chữ thì sửa file này rồi chạy lại.

Mọi con số trong bảng đọc từ `runs/*/result.json` và `metrics_{split}.json` —
không gõ tay số nào, đúng nguyên tắc của scripts/make_tables.py.
"""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

DOC = ROOT / "data_report.docx"
BASE = ROOT / "data_report.base.docx"
FIGS = ROOT / "results" / "figures"

FONT = "Times New Roman"
SIZE = Pt(13)
IMG_W = Inches(6.6)

# ID -> (đầu vào, kiến trúc, vai trò)
ROWS = [
    ("S1", "RGB", "1 luồng", "baseline"),
    ("S2", "IR", "1 luồng", "baseline"),
    ("C1", "RGB → cả hai luồng", "2 luồng, concat+1×1", "đối chứng capacity"),
    ("C2", "IR → cả hai luồng", "2 luồng, concat+1×1", "đối chứng capacity"),
    ("C2b", "IR → cả hai luồng", "2 luồng + cổng SE", "đối chứng của F2b"),
    ("C2c", "IR → cả hai luồng", "2 luồng + cổng modality", "đối chứng của F2c"),
    ("F1", "RGB + IR", "1 luồng, 4 kênh", "fusion sớm"),
    ("F2a", "RGB + IR", "2 luồng, concat+1×1", "fusion giữa"),
    ("F2b", "RGB + IR", "2 luồng + cổng SE", "fusion giữa + attention"),
    ("F2c", "RGB + IR", "2 luồng + cổng modality", "fusion giữa + attention"),
    ("F3", "RGB + IR", "2 × 1 luồng, hợp box", "fusion muộn (WBF)"),
]

DATASETS = [("dronevehicle", "test"), ("vedai", "val")]


# ───────────────────────────────── số liệu ───────────────────────────────────
def runs_of(ds: str, cid: str) -> list[Path]:
    return sorted(d for d in (ROOT / "runs" / ds).glob(f"{cid}_*") if (d / "result.json").exists())


def params(cid: str) -> str:
    """Số tham số, lấy trên DroneVehicle (5 lớp). F3 suy ra từ S1 + S2."""
    if cid == "F3":
        n = sum(json.loads((runs_of("dronevehicle", c)[0] / "result.json")
                           .read_text(encoding="utf-8"))["params"] for c in ("S1", "S2"))
        return f"{n:,}".replace(",", ".") + " *"
    r = runs_of("dronevehicle", cid)
    n = json.loads((r[0] / "result.json").read_text(encoding="utf-8")).get("params")
    return f"{n:,}".replace(",", ".") if n else "—"


def map50(ds: str, split: str, cid: str) -> str:
    vals = []
    for r in runs_of(ds, cid):
        m = r / f"metrics_{split}.json"
        if m.exists():
            vals.append(100 * json.loads(m.read_text(encoding="utf-8"))["all"]["all"]["map50"])
    if not vals:
        return "—"
    v = np.asarray(vals)
    s = f"{v.mean():.2f}".replace(".", ",")
    return s if len(v) == 1 else s + " ± " + f"{v.std():.2f}".replace(".", ",")


# ──────────────────────────────── tiện ích docx ──────────────────────────────
def fmt(run, bold=False, italic=False, size=SIZE, color=None):
    run.font.name = FONT
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return run


def para(doc, text="", bold=False, italic=False, size=SIZE, style=None, align=None,
         space_before=None, space_after=None):
    p = doc.add_paragraph(style=style)
    if text:
        fmt(p.add_run(text), bold=bold, italic=italic, size=size)
    if align is not None:
        p.paragraph_format.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = space_before
    if space_after is not None:
        p.paragraph_format.space_after = space_after
    return p


def rich(doc, parts, style=None):
    """Một đoạn gồm nhiều mẩu (text, bold) — để in đậm giữa câu."""
    p = doc.add_paragraph(style=style)
    for text, bold in parts:
        fmt(p.add_run(text), bold=bold)
    return p


def heading(doc, text):
    return para(doc, text, bold=True, space_before=Pt(12), space_after=Pt(4))


def figure(doc, name, caption):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIGS / name), width=IMG_W)
    c = doc.add_paragraph()
    c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt(c.add_run(caption), italic=True, size=Pt(11), color=RGBColor(0x40, 0x40, 0x40))
    return p


def results_table(doc):
    cols = ["ID", "Đầu vào", "Kiến trúc", "Tham số", "mAP50 DroneVehicle", "mAP50 VEDAI"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for c, name in zip(t.rows[0].cells, cols):
        fmt(c.paragraphs[0].add_run(name), bold=True, size=Pt(11))
    for cid, inp, arch, role in ROWS:
        cells = t.add_row().cells
        vals = [cid, inp, arch, params(cid),
                map50("dronevehicle", "test", cid), map50("vedai", "val", cid)]
        for c, v in zip(cells, vals):
            fmt(c.paragraphs[0].add_run(v), bold=cid in ("F2b", "F2c"), size=Pt(11))
    return t


# ───────────────────────────────── nội dung ──────────────────────────────────
def build(doc):
    heading(doc, "4.1. Attention ở đây là gì")

    para(doc, "Mô hình fusion giữa cơ bản (F2a) cho ảnh RGB đi qua một backbone, ảnh IR đi qua một "
              "backbone khác, rồi ghép hai bộ đặc trưng lại và dùng một lớp Conv 1×1 trộn chúng. "
              "Vấn đề là lớp trộn này học xong thì cố định: mọi ảnh — ban ngày nắng gắt hay ban đêm "
              "tối đen — đều được trộn theo đúng một công thức.")

    rich(doc, [("Attention (cơ chế chú ý) sinh ra để sửa đúng điểm đó. Ta gắn thêm một ", False),
               ("“cổng”", True),
               (" nhỏ: nó nhìn vào chính đặc trưng của ảnh đang xử lý rồi sinh ra một bộ trọng số "
                "riêng cho ảnh đó. Ảnh ban đêm, kênh RGB gần như vô dụng thì cổng hạ trọng số của "
                "RGB xuống và nâng trọng số của IR lên; ảnh ban ngày thì ngược lại. Nói ngắn gọn: "
                "Conv 1×1 là công thức trộn cố định, còn cổng attention là công thức trộn thay đổi "
                "theo từng ảnh.", False)])

    para(doc, "Cách cài đặt (kiểu Squeeze-and-Excitation) gồm ba bước: (1) nén mỗi kênh đặc trưng "
              "thành một con số duy nhất bằng cách lấy trung bình toàn ảnh; (2) cho các con số đó "
              "qua một mạng nơ-ron nhỏ hai lớp để sinh ra trọng số cho từng kênh; (3) nhân trọng số "
              "đó trở lại đặc trưng. Cổng này rất rẻ: chỉ 198.784 tham số, bằng 1,25% so với 15,9 "
              "triệu tham số của F2a — nên nếu nó có ích thì đây là cách cải thiện độ chính xác với "
              "chi phí gần như bằng không.")

    heading(doc, "4.2. Đã cài và thử hai biến thể cổng")

    rich(doc, [("F2b — cổng SE tiêu chuẩn. ", True),
               ("Trọng số của mỗi kênh là một số trong khoảng 0–1 do hàm sigmoid sinh ra, hai "
                "modality hoàn toàn độc lập với nhau.", False)])

    rich(doc, [("F2c — cổng cải tiến. ", True),
               ("Ba thay đổi so với F2b, và điều đáng chú ý là cả ba đều ", False),
               ("không tốn thêm một tham số nào", True), (":", False)])

    para(doc, "1. Khởi tạo đồng nhất. Lúc bắt đầu huấn luyện, cổng cho trọng số đúng bằng 1, nghĩa "
              "là F2c xuất phát y hệt F2a rồi mới học dần cách lệch đi. F2b không có tính chất này: "
              "sigmoid lúc khởi tạo cho khoảng 0,5 nên toàn bộ đặc trưng bị co lại một nửa ngay tại "
              "thời điểm vừa nạp xong trọng số pretrained.", style="List Bullet")

    para(doc, "2. Đọc cả giá trị trung bình lẫn giá trị lớn nhất của mỗi kênh. Một chiếc xe đang "
              "nóng nằm trên nền ruộng lạnh trong ảnh IR là một điểm sáng nhọn; nếu chỉ lấy trung "
              "bình toàn ảnh thì tín hiệu đó bị bôi mất.", style="List Bullet")

    para(doc, "3. Cho hai modality cạnh tranh nhau bằng hàm softmax: trọng số của RGB và của IR tại "
              "cùng một kênh cộng lại bằng hằng số, nên cổng buộc phải chọn tin bên nào nhiều hơn. "
              "Nhờ ràng buộc này mà ta đọc được trực tiếp “mô hình đang tin IR bao nhiêu phần trăm” "
              "cho từng ảnh — chính là số liệu ở Hình B.", style="List Bullet")

    para(doc, "Cả hai biến thể đều có nhóm đối chứng capacity riêng (C2b cho F2b, C2c cho F2c): "
              "cùng kiến trúc và cùng số tham số tuyệt đối, chỉ khác là đưa ảnh IR vào cả hai nhánh "
              "thay vì RGB + IR. Đây là điều kiện bắt buộc để khẳng định phần cải thiện đến từ "
              "thông tin của cảm biến thứ hai chứ không phải từ việc mô hình có nhiều tham số hơn.")

    heading(doc, "4.3. Bảng kết quả")
    results_table(doc)
    para(doc, "Tham số tính trên DroneVehicle (5 lớp). mAP50 đo bằng harness của đề tài. "
              "DroneVehicle: S1, S2, C1, C2, F1, F2a có 2 seed nên ghi kèm độ lệch chuẩn; bốn cấu "
              "hình có cổng mới chạy seed 0. VEDAI: trung bình ± độ lệch chuẩn của 3 fold. "
              "(*) F3 là fusion muộn, ghép kết quả của hai mô hình đã train nên tham số là tổng của "
              "S1 và S2, không phải một mô hình duy nhất.",
          italic=True, size=Pt(11), space_before=Pt(4))

    heading(doc, "4.4. Hình kết quả")
    figure(doc, "fig09_attention_map.png",
           "Hình A. Độ chính xác của các biến thể trên DroneVehicle. Nhóm fusion (xanh lá) tách "
           "hẳn khỏi nhóm đối chứng capacity (cam) — đó là kết quả chính. Nhưng ba biến thể "
           "fusion giữa F2a, F2b, F2c thì nằm sát nhau trong khoảng 0,5 điểm.")
    figure(doc, "fig10_attention_gate.png",
           "Hình B. Cổng attention thật sự học được gì. Trái: tỉ trọng IR giảm dần khi cảnh sáng "
           "lên — đúng hướng giả thuyết. Phải: nhưng toàn bộ 8.980 ảnh nằm trong một dải rộng "
           "0,019, tức cổng gần như là một hằng số.")
    figure(doc, "fig11_attention_curves.png",
           "Hình C. Đường cong huấn luyện. Ba biến thể fusion bám sát nhau và cùng nằm trên "
           "đường đối chứng C2 suốt 60 epoch — khoảng cách giữa chúng nhỏ hơn nhiều so với "
           "khoảng cách tới đối chứng. Cả ba đã bão hoà, không có dấu hiệu thiếu huấn luyện.")
    figure(doc, "fig12_attention_per_class.png",
           "Hình D. AP@50 theo từng lớp trên DroneVehicle. Lớp car chiếm 86% số box và đã bão "
           "hoà ở 97,5% ngay với riêng IR nên fusion gần như không thêm được gì; phần cải thiện "
           "tập trung ở truck (+2,9 điểm) và van (+2,3). Riêng freight_car thì fusion còn kém "
           "hơn đối chứng 0,8 điểm.")

    heading(doc, "4.5. Nhận xét")

    rich(doc, [("1. Kết quả chính của đề tài vẫn vững. ", True),
               ("Fusion thắng đối chứng capacity của chính nó ở cả ba biến thể: F2a − C2 = +9,34 "
                "điểm mAP50, F2b − C2b = +5,80, F2c − C2c = +6,35 (khoảng tin cậy 95% "
                "[+1,21; +11,79], p = 0,007; bootstrap ghép cặp theo ảnh trên VEDAI fold 01). "
                "Nghĩa là phần cải thiện đo được đến từ thông tin bổ sung của cảm biến thứ hai, "
                "không phải từ việc mô hình to gấp đôi.", False)])

    rich(doc, [("2. Nhưng attention không cải thiện thêm được gì so với phép ghép thường. ", True),
               ("Trên DroneVehicle, F2b hơn F2a 0,47 điểm và F2c hơn 0,07 điểm; trên VEDAI thì F2c "
                "lại thấp hơn F2a 0,90 điểm. Mọi chênh lệch này đều nhỏ hơn độ dao động giữa các "
                "lần chạy (VEDAI: ±2,4 điểm giữa các fold), nên không thể kết luận biến thể nào "
                "tốt hơn.", False)])

    rich(doc, [("3. Hình B cho biết vì sao. ", True),
               ("Cổng học ra một thiên lệch gần như cố định — nghiêng về IR ở mức 0,56 — thay vì "
                "thay đổi theo từng ảnh như thiết kế mong đợi. Toàn bộ 8.980 ảnh test chỉ trải trên "
                "một dải rộng 0,019. Một hệ số nhân gần như hằng số thì lớp Conv 1×1 ngay sau đó "
                "học bù lại được, nên cổng không đóng góp gì thêm. Trên VEDAI, nơi tập huấn luyện "
                "chỉ có 1.089 ảnh, cổng thậm chí đứng yên ở 0,50 — không học được gì.", False)])

    rich(doc, [("4. Tuy vậy cơ chế thì đúng hướng, chỉ là quá yếu. ", True),
               ("Tương quan giữa tỉ trọng IR và độ sáng của cảnh là −0,72 tại mức hợp nhất P4: cảnh "
                "càng tối, cổng càng nghiêng về IR, đúng như giả thuyết H3. Quan hệ này đơn điệu "
                "gần như hoàn hảo qua cả 10 decile chiếu sáng. Vấn đề duy nhất là biên độ của nó "
                "nhỏ hơn khoảng 100 lần so với mức cần thiết để làm thay đổi đặc trưng đi vào đầu "
                "dò.", False)])

    rich(doc, [("5. Phép đo đã được kiểm chứng. ", True),
               ("Ở cấu hình đối chứng C2c, khi cả hai nhánh cùng nhận ảnh IR và do đó không có gì "
                "để lựa chọn, cổng đứng yên ở 0,498–0,508. Điều này xác nhận tỉ trọng đo được ở F2c "
                "là phản ứng thật với sự khác biệt giữa hai modality, không phải hiện vật của cách "
                "đo.", False)])

    rich(doc, [("6. Fusion giúp ở đâu thì rõ hơn là giúp bao nhiêu. ", True),
               ("Nhìn theo lớp (Hình D), lợi ích không rải đều. Lớp car chiếm 86% số box và đã "
                "đạt 97,5% ngay với IR đơn thuần nên không còn chỗ để cải thiện, mà chính lớp "
                "này lại chi phối mAP tổng thể; toàn bộ phần chênh lệch thực chất đến từ truck "
                "và van. Trên VEDAI (không vẽ hình ở đây) quy luật còn rõ hơn: lợi ích dồn vào "
                "các lớp hiếm — other +21,5 điểm, tractor +15,2, boat +11,2 — tức những lớp mà "
                "một modality đơn lẻ khó phân biệt, dù mỗi lớp chỉ có 17–25 box nên con số dao "
                "động mạnh. Đây là câu trả lời có ích hơn cho câu hỏi “khi nào nên dùng fusion” "
                "so với một con số mAP tổng.", False)])

    rich(doc, [("Kết luận. ", True),
               ("Attention là một ý tưởng hợp lý về mặt lý thuyết và rất rẻ về mặt tham số, nhưng "
                "trong thí nghiệm này nó không cải thiện được độ chính xác. Đóng góp của phần này "
                "không nằm ở con số mAP mà ở chỗ giải thích được nguyên nhân bằng số liệu đo trực "
                "tiếp từ bên trong mô hình: cổng hội tụ về một thiên lệch tĩnh chứ không phải một "
                "phép lựa chọn động. Đây là kết quả âm tính có cơ chế, và nó đúng với tinh thần "
                "phương pháp luận của đề tài — báo cáo cả những gì không hiệu quả, kèm bằng chứng "
                "vì sao.", False)])


def main():
    if not BASE.exists():
        shutil.copy2(DOC, BASE)
        print(f"  sao luu ban goc -> {BASE.name}")
    doc = Document(str(BASE))
    build(doc)
    doc.save(str(DOC))
    print(f"  -> {DOC.name}: {len(doc.paragraphs)} doan, {len(doc.tables)} bang, "
          f"{len(doc.inline_shapes)} hinh")


if __name__ == "__main__":
    main()
